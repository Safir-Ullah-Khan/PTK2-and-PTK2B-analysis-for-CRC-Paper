#!/usr/bin/env python3
import subprocess, sys
for pkg in ['GEOparse']:
    try:
        __import__(pkg)
    except ImportError:
        subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', pkg], check=True)

import os, re, gzip, urllib.request, warnings
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from scipy import stats
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')
try:
    from google.colab import drive
    drive.mount('/content/drive')
except Exception:
    pass

BASE     = '/content/drive/MyDrive/FAK_PYK2_revision'
DATA_DIR = os.path.join(BASE, 'data')
OUT      = os.path.join(BASE, 'Script_PYK2_keratin_complete')
for d in (DATA_DIR, OUT):
    os.makedirs(d, exist_ok=True)
XLSX = os.path.join(OUT, 'Supplementary_Tables_PYK2_keratin_HGSOC.xlsx')
INCLUDE_GENOMEWIDE = False

COHORTS = {
    'TCGA_RNAseq':        dict(label='TCGA',      src='xena',
                               expr='TCGA.OV.sampleMap%2FHiSeqV2.gz',
                               ref='TCGA Research Network, Nature 2011', acc='TCGA-OV',
                               platform='Illumina HiSeq RNA-seq (RSEM, log2)',
                               source='UCSC Xena (TCGA hub)'),
    'Tothill_GSE9891':    dict(label='Tothill',    src='geo',  gse='GSE9891',
                               ref='Tothill et al., Clin Cancer Res 2008', acc='GSE9891',
                               platform='Affymetrix HG-U133 Plus 2.0',
                               source='GEO series matrix'),
    'Bonome_GSE26712':    dict(label='Bonome',     src='geo',  gse='GSE26712',
                               ref='Bonome et al., Cancer Res 2008', acc='GSE26712',
                               platform='Affymetrix HG-U133A',
                               source='GEO series matrix'),
    'Yoshihara_GSE32062': dict(label='Yoshihara',  src='geo',  gse='GSE32062',
                               gpls=('GPL6480',),
                               ref='Yoshihara et al., Clin Cancer Res 2012', acc='GSE32062',
                               platform='Agilent 4x44K (GPL6480)',
                               source='GEO series matrix'),
    'Konecny_GSE53963':   dict(label='Konecny',    src='geo',  gse='GSE53963',
                               ref='Konecny et al., J Natl Cancer Inst 2014', acc='GSE53963',
                               platform='Agilent 4x44K',
                               source='GEO series matrix (TCGA samples removed)'),
}
LAB  = {k: v['label'] for k, v in COHORTS.items()}
MARK = dict(zip(COHORTS, ['o', 's', '^', 'v', 'D']))

XENA       = 'https://tcga-xena-hub.s3.us-east-1.amazonaws.com/download/'
TCGA_CLIN  = 'TCGA.OV.sampleMap%2FOV_clinicalMatrix'
ABSOLUTE_URL = 'https://api.gdc.cancer.gov/data/4f277128-f793-4354-a13d-30cc7fe9f6b5'

STROMAL = ['COL1A1', 'COL1A2', 'COL3A1', 'COL5A1', 'DCN', 'LUM', 'FAP', 'PDGFRB', 'THY1', 'FBN1']
IMMUNE  = ['PTPRC', 'CD2', 'CD3E', 'CD53', 'LAPTM5', 'CORO1A', 'TYROBP', 'FCER1G', 'LCP2', 'CD48']
KER     = ['KRT7', 'KRT8', 'KRT18', 'KRT19', 'KRT80', 'KRT14']
SCORE   = 'Keratin programme'

G1S = ['MCM5', 'PCNA', 'TYMS', 'FEN1', 'MCM2', 'MCM4', 'RRM1', 'UNG', 'GINS2', 'MCM6', 'CDCA7', 'DTL',
       'PRIM1', 'UHRF1', 'CENPU', 'HELLS', 'RFC2', 'RPA2', 'NASP', 'RAD51AP1', 'GMNN', 'WDR76', 'SLBP',
       'CCNE2', 'UBR7', 'POLD3', 'MSH2', 'ATAD2', 'RAD51', 'RRM2', 'CDC45', 'CDC6', 'EXO1', 'TIPIN',
       'DSCC1', 'BLM', 'CASP8AP2', 'USP1', 'CLSPN', 'POLA1', 'CHAF1B', 'BRIP1', 'E2F8']
G2M = ['HMGB2', 'CDK1', 'NUSAP1', 'UBE2C', 'BIRC5', 'TPX2', 'TOP2A', 'NDC80', 'CKS2', 'NUF2', 'CKS1B',
       'MKI67', 'TMPO', 'CENPF', 'TACC3', 'PIMREG', 'SMC4', 'CCNB2', 'CKAP2L', 'CKAP2', 'AURKB', 'BUB1',
       'KIF11', 'ANP32E', 'TUBB4B', 'GTSE1', 'KIF20B', 'HJURP', 'CDCA3', 'JPT1', 'CDC20', 'TTK', 'CDC25C',
       'KIF2C', 'RANGAP1', 'NCAPD2', 'DLGAP5', 'CDCA2', 'CDCA8', 'ECT2', 'KIF23', 'HMMR', 'AURKA', 'PSRC1',
       'ANLN', 'LBR', 'CKAP5', 'CENPE', 'CTCF', 'NEK2', 'G2E3', 'GAS2L3', 'CBX5', 'CENPA']
ALIAS  = {'MLF1IP': 'CENPU', 'FAM64A': 'PIMREG', 'HN1': 'JPT1'}
PROLIF = G1S + G2M

Q_HIGH, EXPR_PCTL, EXPR_FRAC, MIN_COHORTS = 0.75, 25, 0.10, 3
MODELS = {'adjusted': ['stromal', 'immune'],
          'adjusted + proliferation': ['stromal', 'immune', 'Proliferation']}
MN = {'adjusted': 'Stroma + immune',
      'adjusted + proliferation': 'Stroma + immune + proliferation'}

BLACK, RED, CURVE, GREY = '#000000', '#B2182B', '#BDBDBD', '#7F7F7F'
plt.rcParams.update({
    'font.family': 'sans-serif', 'font.sans-serif': ['Arial', 'Liberation Sans', 'DejaVu Sans'],
    'font.size': 11, 'font.weight': 'bold', 'axes.titlesize': 11, 'axes.titleweight': 'bold',
    'axes.labelsize': 11, 'axes.labelweight': 'bold', 'xtick.labelsize': 11, 'ytick.labelsize': 11,
    'legend.fontsize': 10, 'text.color': BLACK, 'axes.labelcolor': BLACK, 'axes.edgecolor': BLACK,
    'xtick.color': BLACK, 'ytick.color': BLACK, 'axes.linewidth': 1.0,
    'axes.spines.top': False, 'axes.spines.right': False,
    'pdf.fonttype': 42, 'ps.fonttype': 42, 'svg.fonttype': 'none',
})

def save_fig(fig, name):
    for ext, kw in [('pdf', {}), ('svg', {}), ('png', {'dpi': 600}),
                    ('tiff', {'dpi': 600, 'pil_kwargs': {'compression': 'tiff_lzw'}})]:
        fig.savefig(os.path.join(OUT, f'{name}.{ext}'), bbox_inches='tight', facecolor='white', **kw)
    plt.show(); plt.close(fig)

EXCLUDE_VAL  = re.compile(r'normal|\bose\b|surface epith|\blmp\b|low malignant|borderline|benign|cell line', re.I)
SEROUS_VAL   = re.compile(r'ser', re.I)
LOWGRADE_VAL = re.compile(r'^\s*(1|g1|grade\s*1|i|low|well)\b', re.I)

def download(url, path):
    if os.path.exists(path) and os.path.getsize(path) > 1000:
        return path
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=300) as r, open(path, 'wb') as f:
        f.write(r.read())
    return path

def to_log2(df):
    v = df.values[np.isfinite(df.values)]
    return np.log2(df.clip(lower=0) + 1) if np.nanpercentile(v, 99) > 50 else df

def collapse_probes(expr, symbols):
    s = symbols.reindex(expr.index).astype(str).str.split(r'\s*///\s*|\s*//\s*').str[0].str.strip()
    ok = s.notna() & (s != '') & (s != 'nan') & (s != '---')
    e = expr[ok].copy(); e['__sym'] = s[ok]; e['__mean'] = e.drop(columns='__sym').mean(1)
    e = e.sort_values('__mean', ascending=False).drop_duplicates('__sym')
    return e.set_index('__sym').drop(columns='__mean')

def load_tcga(name, cfg):
    log = []
    p = download(XENA + cfg['expr'], os.path.join(DATA_DIR, cfg['expr'].replace('%2F', '_')))
    expr = pd.read_csv(p, sep='\t', index_col=0, compression='gzip')
    expr.columns = [c[:15] for c in expr.columns]
    expr = expr.loc[:, ~expr.columns.duplicated()]
    log.append(('All samples', expr.shape[1]))
    cp = download(XENA + TCGA_CLIN, os.path.join(DATA_DIR, 'TCGA_OV_clinicalMatrix.tsv'))
    clin = pd.read_csv(cp, sep='\t', index_col=0)
    clin.index = [i[:15] for i in clin.index]
    clin = clin[~clin.index.duplicated()]
    keep = [c for c in expr.columns if c[13:15] == '01']
    log.append(('Primary tumour (sample type 01)', len(keep)))
    meta = clin.reindex(keep)
    hist  = meta.get('histological_type', pd.Series(index=keep, dtype=object)).astype(str)
    grade = meta.get('neoplasm_histologic_grade', pd.Series(index=keep, dtype=object)).astype(str)
    serous = hist.str.contains('serous', case=False, na=False) | hist.isin(['nan', 'None'])
    keep = [k for k in keep if serous[k]]
    log.append(('Serous histology (or not recorded)', len(keep)))
    lowg = grade.str.upper().isin(['G1', 'GB'])
    keep = [k for k in keep if not lowg[k]]
    log.append(('Not low grade / borderline (G1, GB excluded)', len(keep)))
    pts = pd.Series(keep, index=keep).str[:12]
    keep = list(pts[~pts.duplicated()].index)
    log.append(('One sample per patient', len(keep)))
    meta = meta.reindex(keep)
    out = pd.DataFrame(index=keep)
    out['age']      = pd.to_numeric(meta.get('age_at_initial_pathologic_diagnosis'), errors='coerce')
    out['stage']    = meta.get('clinical_stage')
    out['grade']    = meta.get('neoplasm_histologic_grade')
    out['OS_time']  = pd.to_numeric(meta.get('OS.time', meta.get('_OS')), errors='coerce')
    out['OS_event'] = pd.to_numeric(meta.get('OS', meta.get('_OS_IND')), errors='coerce')
    try:
        ap = download(ABSOLUTE_URL, os.path.join(DATA_DIR, 'TCGA_mastercalls.abs_tables_JSedit.fixed.txt'))
        ab = pd.read_csv(ap, sep='\t'); ab.index = ab['array'].astype(str).str[:15]
        out['purity_ABSOLUTE'] = pd.to_numeric(ab['purity'], errors='coerce').reindex(keep)
    except Exception as e:
        print(f'   ABSOLUTE purity not available ({type(e).__name__})')
    return to_log2(expr[keep]), out, log

def read_series_matrix(gse, gpls=()):
    path  = os.path.join(DATA_DIR, f'{gse}_series_matrix.txt.gz')
    base  = f'https://ftp.ncbi.nlm.nih.gov/geo/series/{gse[:-3]}nnn/{gse}/matrix/'
    names = [f'{gse}_series_matrix.txt.gz'] + [f'{gse}-{g}_series_matrix.txt.gz' for g in gpls]
    last  = None
    for nm in names:
        try:
            download(base + nm, path); last = None; break
        except Exception as e:
            last = e
    if last is not None:
        raise last
    meta_rows, table, in_table = {}, [], False
    with gzip.open(path, 'rt', errors='replace') as fh:
        for line in fh:
            line = line.rstrip('\n')
            if line.startswith('!series_matrix_table_begin'):
                in_table = True; continue
            if line.startswith('!series_matrix_table_end'):
                break
            if in_table:
                table.append(line.split('\t')); continue
            if line.startswith('!Sample_'):
                parts = [x.strip('"') for x in line.split('\t')]
                meta_rows.setdefault(parts[0], []).append(parts[1:])
    hdr  = [h.strip('"') for h in table[0]]
    expr = pd.DataFrame([r[1:] for r in table[1:]],
                        index=[r[0].strip('"') for r in table[1:]],
                        columns=hdr[1:]).apply(pd.to_numeric, errors='coerce')
    gsm  = meta_rows['!Sample_geo_accession'][0]
    info = pd.DataFrame(index=gsm)
    for key in ['!Sample_title', '!Sample_source_name_ch1', '!Sample_platform_id']:
        if key in meta_rows:
            info[key.replace('!Sample_', '')] = meta_rows[key][0]
    for key in [k for k in meta_rows if k.startswith('!Sample_characteristics')]:
        for vals in meta_rows[key]:
            for g, v in zip(gsm, vals):
                if ':' in v:
                    k, val = v.split(':', 1)
                    info.loc[g, k.strip().lower()] = val.strip()
    return expr, info

def load_geo(name, cfg):
    import GEOparse
    log = []
    expr, info = read_series_matrix(cfg['gse'], cfg.get('gpls', ()))
    log.append(('All samples', expr.shape[1]))
    gpl_id = info['platform_id'].mode()[0]
    gpl    = GEOparse.get_GEO(geo=gpl_id, destdir=DATA_DIR, silent=True)
    tab    = gpl.table.copy(); tab['ID'] = tab['ID'].astype(str); tab = tab.set_index('ID')
    sym_col = next((c for c in ['Gene Symbol', 'GENE_SYMBOL', 'Symbol', 'SYMBOL',
                                'gene_symbol', 'Gene symbol', 'ORF'] if c in tab.columns), None)
    if sym_col is None:
        raise KeyError(f'{name}: no gene-symbol column in {gpl_id}: {list(tab.columns)[:15]}')
    expr = collapse_probes(expr, tab[sym_col])
    keep = pd.Series(True, index=info.index)
    type_keys  = [c for c in info.columns if re.search(r'type|tissue|source|malignan|title', c, re.I)]
    hist_keys  = [c for c in info.columns if re.search(r'subtype|histolog|morpholog|histotype', c, re.I)]
    grade_keys = [c for c in info.columns if re.search(r'grade', c, re.I)]
    excl = info[type_keys].astype(str).apply(
        lambda r: bool(EXCLUDE_VAL.search(' | '.join(r))), axis=1) \
        if type_keys else pd.Series(False, index=info.index)
    keep &= ~excl
    log.append((f'Excluding normal / LMP / borderline / cell lines (fields: {type_keys})', int(keep.sum())))
    if hist_keys:
        h = info[hist_keys].astype(str).agg(' | '.join, axis=1)
        recorded = ~h.str.fullmatch(r'(nan|None|\s|\|)*')
        keep &= ~recorded | h.str.contains(SEROUS_VAL)
        log.append((f'Serous where histology recorded (fields: {hist_keys})', int(keep.sum())))
    if grade_keys:
        g = info[grade_keys[0]].astype(str)
        keep &= ~g.str.match(LOWGRADE_VAL)
        log.append((f'Excluding grade 1 / low grade (field: {grade_keys[0]})', int(keep.sum())))
    tcga_keys = [c for c in info.columns if re.search(r'tcga', c, re.I)]
    if tcga_keys:
        is_tcga = info[tcga_keys].astype(str).apply(
            lambda r: bool(re.search(r'TCGA-', ' '.join(r))), axis=1)
        keep &= ~is_tcga
        log.append((f'Removing samples also in TCGA (field: {tcga_keys[0]})', int(keep.sum())))
    ids  = [s for s in info.index[keep] if s in expr.columns]
    meta = info.loc[ids].drop(columns=[c for c in ['platform_id'] if c in info.columns])
    return to_log2(expr[ids]), meta, log

def load_cohort(name, cfg):
    ce = os.path.join(DATA_DIR, f'{name}_v3_harmonised_expr.tsv.gz')
    cm = os.path.join(DATA_DIR, f'{name}_v3_harmonised_meta.tsv')
    cl = os.path.join(DATA_DIR, f'{name}_v3_filter_log.tsv')
    if all(os.path.exists(x) for x in (ce, cm, cl)):
        print(f'   {name}: loaded from cache')
        return (pd.read_csv(ce, sep='\t', index_col=0),
                pd.read_csv(cm, sep='\t', index_col=0),
                [tuple(r) for r in pd.read_csv(cl, sep='\t').values])
    print(f'   {name}: downloading and harmonising ...')
    expr, meta, log = (load_tcga if cfg['src'] == 'xena' else load_geo)(name, cfg)
    expr = expr.loc[~expr.index.duplicated()]
    expr.to_csv(ce, sep='\t', compression='gzip')
    meta.to_csv(cm, sep='\t')
    pd.DataFrame(log, columns=['step', 'n']).to_csv(cl, sep='\t', index=False)
    print(f'   {name}: saved to cache ({expr.shape[1]} tumours, {expr.shape[0]:,} genes)')
    return expr, meta, log

def is_noise(g):
    return bool(re.match(r'^(MT-|RPL\d|RPS\d|MRPL\d|MRPS\d)', str(g)))

def expressed_genes(expr):
    thr = np.percentile(expr.mean(1), EXPR_PCTL)
    return set(expr.index[(expr > thr).mean(1) >= EXPR_FRAC])

def zmean(e, genes):
    g = [x for x in genes if x in e.index]
    return e.loc[g].sub(e.loc[g].mean(1), axis=0).div(e.loc[g].std(1).replace(0, 1), axis=0).mean(0)

def rank_score(expr, genes):
    R = expr.rank(axis=0) / len(expr)
    return R.loc[genes].mean(0) - 0.5

def resid_ranks(d, cols, covs):
    R = d[cols + covs].dropna().rank()
    A = np.column_stack([np.ones(len(R)), R[covs].values])
    return pd.DataFrame(
        {c: R[c].values - A @ np.linalg.lstsq(A, R[c].values, rcond=None)[0]
         for c in cols}, index=R.index)

def residualise(R, covs):
    A = np.column_stack([np.ones(R.shape[1]), covs])
    return R - (A @ np.linalg.lstsq(A, R.T, rcond=None)[0]).T

def hedges(a, b):
    na, nb = a.shape[1], b.shape[1]
    sp = np.sqrt(((na - 1) * a.var(1, ddof=1) + (nb - 1) * b.var(1, ddof=1)) / (na + nb - 2))
    sp = np.where(sp == 0, np.nan, sp)
    g  = (a.mean(1) - b.mean(1)) / sp * (1 - 3 / (4 * (na + nb) - 9))
    return g, (na + nb) / (na * nb) + g ** 2 / (2 * (na + nb))

def dl_rows(G, V):
    W  = np.where(np.isfinite(G) & np.isfinite(V), 1 / V, 0.0)
    Gz = np.nan_to_num(G)
    k  = (W > 0).sum(1); sw = W.sum(1); sw[sw == 0] = np.nan
    ybar = (W * Gz).sum(1) / sw
    Q    = (W * (Gz - ybar[:, None]) ** 2).sum(1)
    c    = sw - (W ** 2).sum(1) / sw
    tau2 = np.clip((Q - (k - 1)) / c, 0, None)
    Vv   = np.where(W > 0, 1 / np.where(W > 0, W, 1), np.inf)
    Ws   = np.where(W > 0, 1 / (Vv + tau2[:, None]), 0.0)
    est  = (Ws * Gz).sum(1) / Ws.sum(1)
    se   = np.sqrt(1 / Ws.sum(1))
    I2   = np.where(Q > 0, np.clip((Q - (k - 1)) / Q, 0, None) * 100, 0)
    return est, se, k, I2

def bh(p):
    p = np.asarray(p, float); n = len(p)
    o = np.argsort(p); r = np.empty(n); r[o] = np.arange(1, n + 1)
    q = np.minimum.accumulate((p * n / r)[o][::-1])[::-1]
    out = np.empty(n); out[o] = np.minimum(q, 1); return out

def stars(q):
    return '***' if q < 0.001 else '**' if q < 0.01 else '*' if q < 0.05 else 'ns'

print('\n' + '=' * 70)
print('PART 1 — downloading / loading cohorts')
print('=' * 70)

gw, rows, used_rows, cohort_rows, step_rows, expressed_count = {}, [], [], [], [], {}

for name, cfg in COHORTS.items():
    print(f'\n>>> {name}')
    try:
        raw_expr, meta_df, log = load_cohort(name, cfg)
    except Exception as e:
        print(f'   FAILED: {type(e).__name__}: {e}'); continue
    expr = raw_expr.apply(pd.to_numeric, errors='coerce').replace([np.inf, -np.inf], np.nan)
    expr = expr.rename(index=ALIAS)
    expr = expr.loc[~expr.index.duplicated()]
    expr = expr.loc[:, expr.notna().mean(0) > 0.5]
    expr = expr[expr.isna().mean(1) <= 0.20]
    expr = expr.T.fillna(expr.median(1)).T
    expd = {g for g in expressed_genes(expr) if not is_noise(g)}
    for g in expd:
        expressed_count[g] = expressed_count.get(g, 0) + 1
    print(f'   {cfg["label"]}: {expr.shape[1]} tumours; {len(expd):,} expressed genes')
    for i, (step, n) in enumerate(log):
        fld = re.search(r'\(fields?:\s*(.*)\)$', str(step))
        step_rows.append({'Cohort': cfg['label'], 'Step': i + 1,
                          'Inclusion step': re.sub(r'\s*\(fields?:.*\)$', '', str(step)),
                          'Annotation fields used': fld.group(1) if fld else '',
                          'Tumours remaining': int(n)})
    step_rows.append({'Cohort': cfg['label'], 'Step': len(log) + 1,
                      'Inclusion step': 'Analysed (tumours with < 50% missing values)',
                      'Annotation fields used': '', 'Tumours remaining': expr.shape[1]})
    age = np.nan
    ac  = [c for c in meta_df.columns if re.fullmatch(r'age.*', str(c), re.I)]
    if ac:
        a   = pd.to_numeric(meta_df[ac[0]], errors='coerce')
        age = float(a.median()) if a.notna().sum() > 10 else np.nan
    d = pd.DataFrame(index=expr.columns)
    d['PTK2B']   = expr.loc['PTK2B']
    d['stromal'] = zmean(expr, STROMAL)
    d['immune']  = zmean(expr, IMMUNE)
    ker = [g for g in KER if g in expd]
    for g in ker:
        d[g] = expr.loc[g]
    for s, gs in [(SCORE, KER), ('Proliferation', PROLIF)]:
        ok = [g for g in gs if g in expd]
        used_rows.append({'Cohort': cfg['label'], 'Gene set': s,
                          'Genes defined': len(gs),
                          'Genes passing expression filter': len(ok),
                          'Genes used': ', '.join(ok)})
        if len(ok) >= max(3, 0.6 * len(gs)):
            d[s] = rank_score(expr, ok)
    E    = expr.loc[sorted(expd | {'PTK2B'})]
    covs = np.column_stack([d.stromal.rank().values, d.immune.rank().values])
    X    = residualise(E.rank(axis=1).values, covs)
    k_   = X[list(E.index).index('PTK2B')]
    hi   = k_ > np.quantile(k_, Q_HIGH)
    lo   = k_ <= np.quantile(k_, 1 - Q_HIGH)
    g_, v_ = hedges(X[:, hi], X[:, lo])
    gw[name] = pd.DataFrame({'g': g_, 'v': v_}, index=E.index)
    feats = ker + ([SCORE] if SCORE in d else [])
    for model, cv in MODELS.items():
        if 'Proliferation' in cv and 'Proliferation' not in d:
            continue
        A   = resid_ranks(d, ['PTK2B'] + feats, cv)
        hiA = A.PTK2B > A.PTK2B.quantile(Q_HIGH)
        loA = A.PTK2B <= A.PTK2B.quantile(1 - Q_HIGH)
        gg, vv = hedges(A.loc[hiA, feats].values.T, A.loc[loA, feats].values.T)
        for f, a__, b__ in zip(feats, gg, vv):
            rows.append({'cohort': name, 'model': model, 'feature': f,
                         'n_high': int(hiA.sum()), 'n_low': int(loA.sum()),
                         'g': a__, 'var_g': b__})
    cohort_rows.append({
        'Cohort': cfg['label'], 'Reference': cfg['ref'], 'Accession': cfg['acc'],
        'Platform': cfg['platform'], 'Data source': cfg['source'],
        'Samples in source': int(log[0][1]) if log else np.nan,
        'HGSOC after inclusion criteria': int(log[-1][1]) if log else np.nan,
        'Tumours analysed': expr.shape[1],
        'PTK2B-high (top quartile)': int(hi.sum()),
        'PTK2B-low (bottom quartile)': int(lo.sum()),
        'Median age, years': round(age, 1) if np.isfinite(age) else 'Not recorded'})

print('\n' + '=' * 70)
print('PART 2 — pooling and figures')
print('=' * 70)

universe  = sorted(g for g, n in expressed_count.items() if n >= MIN_COHORTS)
G_mat     = pd.concat({c: gw[c]['g'] for c in gw}, axis=1).reindex(universe)
V_mat     = pd.concat({c: gw[c]['v'] for c in gw}, axis=1).reindex(universe)
est, se, kk, I2 = dl_rows(G_mat.values, V_mat.values)
de = pd.DataFrame({'pooled_g': est, 'se': se, 'z': est / se, 'k': kk, 'I2': I2}, index=universe)
de = de[(de.k >= MIN_COHORTS) & np.isfinite(de.z)]
de = de.drop(index=[x for x in STROMAL + IMMUNE + ['PTK2', 'PTK2B'] if x in de.index])
de['p']   = 2 * stats.norm.sf(np.abs(de.z))
de['FDR'] = bh(de.p.values)
de        = de.sort_values('z', ascending=False)
de['rank'] = np.arange(1, len(de) + 1)
de.to_csv(os.path.join(OUT, 'genomewide_PTK2B_high_vs_low_adjusted.csv'))

per  = pd.DataFrame(rows)
meta = []
for (model, f), sub in per.groupby(['model', 'feature']):
    if len(sub) < MIN_COHORTS:
        continue
    e1, s1, k1, i1 = dl_rows(sub.g.values[None, :], sub.var_g.values[None, :])
    meta.append({'model': model, 'feature': f, 'k': int(k1[0]),
                 'n_high': int(sub.n_high.sum()), 'n_low': int(sub.n_low.sum()),
                 'pooled_g': e1[0], 'lo': e1[0] - 1.96 * s1[0],
                 'hi': e1[0] + 1.96 * s1[0],
                 'p': 2 * stats.norm.sf(abs(e1[0] / s1[0])), 'I2': i1[0]})
meta = pd.DataFrame(meta)
meta['FDR'] = np.nan
for _, idx in meta.groupby('model').groups.items():
    meta.loc[idx, 'FDR'] = bh(meta.loc[idx, 'p'].values)
per.to_csv(os.path.join(OUT, 'keratin_per_cohort.csv'), index=False)
meta.to_csv(os.path.join(OUT, 'keratin_pooled.csv'), index=False)
cohorts_run = [c for c in COHORTS if c in gw]

N   = len(de)
hit = de.loc[[g for g in KER if g in de.index]].sort_values('rank')
fig, ax = plt.subplots(figsize=(5.6, 4.6))
ax.plot(de['rank'], de['z'], color=CURVE, lw=2.4, zorder=1)
ax.axhline(0, color=BLACK, lw=0.8, zorder=1)
ax.scatter(hit['rank'], hit['z'], s=46, color=RED, edgecolor=BLACK, lw=0.7, zorder=3)
zmax, zmin = de.z.max(), de.z.min(); yr = zmax - zmin
ax.set_xlim(-0.02 * N, 1.02 * N)
ax.set_ylim(zmin - 0.06 * yr, zmax + 0.06 * yr)
fig.canvas.draw()
gap = 14.0 / (ax.bbox.height * 72 / fig.dpi / (ax.get_ylim()[1] - ax.get_ylim()[0]))
ys, prev = np.array(hit.z, float), -np.inf
for i in np.argsort(ys):
    ys[i] = max(ys[i], prev + gap); prev = ys[i]
x_lab = min(hit['rank'].max() + 0.08 * N, 0.55 * N)
for (g, r), yl in zip(hit.iterrows(), ys):
    ax.annotate(g, xy=(r['rank'], r.z), xytext=(x_lab, yl), textcoords='data',
                ha='left', va='center', fontsize=11, fontstyle='italic',
                fontweight='bold', color=BLACK,
                arrowprops=dict(arrowstyle='-', color=RED, lw=1.0, shrinkA=2, shrinkB=3, relpos=(0, 0.5)))
ax.set_ylim(zmin - 0.06 * yr, max(zmax + 0.06 * yr, ys.max() + gap))
ax.set_xlabel('Genes ranked by pooled z')
ax.set_ylabel('Pooled z (PTK2B high vs low)')
ax.set_title('Keratin programme in PYK2-high tumours', loc='left')
ax.text(0.97, 0.04, 'Stromal/immune-adjusted', transform=ax.transAxes, ha='right', va='bottom', fontsize=10)
plt.tight_layout(); save_fig(fig, 'Fig1_PYK2_ranked_keratin_genes')

pA = per[per.model == 'adjusted']
mA = meta[meta.model == 'adjusted'].set_index('feature')
genes = [g for g in mA.loc[[x for x in KER if x in mA.index]].sort_values('pooled_g', ascending=False).index]
feats = genes + [SCORE]
xpos  = {f: i for i, f in enumerate(genes)}
xpos[SCORE] = len(genes) + 0.6
off   = dict(zip(cohorts_run, np.linspace(-0.27, 0.27, len(cohorts_run))))
fig, ax = plt.subplots(figsize=(7.6, 4.4))
ax.axhline(0, color=BLACK, lw=0.9, ls='--', zorder=1)
for _, r in pA.iterrows():
    ax.scatter(xpos[r.feature] + off[r.cohort], r.g, marker=MARK[r.cohort],
               s=34, facecolor='white', edgecolor=BLACK, lw=1.0, zorder=2)
ytop = []
for f in feats:
    m = mA.loc[f]
    ax.errorbar(xpos[f] + 0.42, m.pooled_g,
                yerr=[[m.pooled_g - m.lo], [m.hi - m.pooled_g]],
                fmt='D', ms=8.5, color=RED, mec=BLACK, mew=0.8,
                ecolor=RED, elinewidth=1.8, capsize=3.5, zorder=3)
    top = max(m.hi, pA[pA.feature == f].g.max())
    ax.text(xpos[f] + 0.08, top + 0.12, stars(m.FDR), ha='center', va='bottom', fontsize=11, fontweight='bold')
    ytop.append(top + 0.45)
ax.axvline(len(genes) - 0.2, color=GREY, lw=0.8)
ax.set_xticks([xpos[f] + 0.08 for f in feats])
ax.set_xticklabels(genes + ['Keratin\nprogramme'])
for t, f in zip(ax.get_xticklabels(), feats):
    if f != SCORE: t.set_fontstyle('italic')
ax.set_xlim(-0.6, xpos[SCORE] + 0.8)
ax.set_ylim(min(-0.2, pA.g.min() - 0.15), max(ytop))
ax.set_ylabel("Hedges' g (PTK2B high vs low)")
ax.set_title('Keratin programme in PYK2-high tumours, per dataset', loc='left')
nn = pA.groupby('cohort')[['n_high', 'n_low']].max()
handles = [Line2D([0], [0], marker=MARK[c], ls='', markerfacecolor='white',
                  markeredgecolor=BLACK, mew=1.0, markersize=7,
                  label=f'{LAB[c]} ({nn.loc[c, "n_high"]} vs {nn.loc[c, "n_low"]})')
           for c in cohorts_run]
handles.append(Line2D([0], [0], marker='D', ls='', markerfacecolor=RED,
                      markeredgecolor=BLACK, markersize=8, label='Pooled (95% CI)'))
ax.legend(handles=handles, frameon=False, loc='upper left',
          bbox_to_anchor=(1.01, 1.0), fontsize=10,
          title='Dataset (high vs low, n)',
          title_fontproperties={'weight': 'bold', 'size': 10})
fig.text(0.0, -0.04,
         'Top vs bottom quartile of PTK2B; stromal/immune-adjusted; random-effects pooling; BH-FDR: * < 0.05, ** < 0.01, *** < 0.001.',
         ha='left', fontsize=9.5)
plt.tight_layout(); save_fig(fig, 'Fig2_PYK2_keratin_genes_per_dataset')

fig, ax = plt.subplots(figsize=(8.6, 4.8))
ytop = []
for f in feats:
    for model, dx, face in [('adjusted', -0.17, RED), ('adjusted + proliferation', 0.17, 'white')]:
        m = meta[(meta.model == model) & (meta.feature == f)]
        if m.empty: continue
        m = m.iloc[0]
        ax.errorbar(xpos[f] + dx, m.pooled_g,
                    yerr=[[m.pooled_g - m.lo], [m.hi - m.pooled_g]],
                    fmt='D', ms=8.5, mfc=face, mec=RED, mew=1.6,
                    ecolor=RED, elinewidth=1.8, capsize=3.5, zorder=3)
        ax.text(xpos[f] + dx, m.hi + 0.04, stars(m.FDR), ha='center', va='bottom', fontsize=10, fontweight='bold')
        ytop.append(m.hi + 0.3)
ax.axhline(0, color=BLACK, lw=0.9, ls='--', zorder=1)
ax.axvline(len(genes) - 0.2, color=GREY, lw=0.8)
ax.set_xticks([xpos[f] for f in feats])
ax.set_xticklabels(genes + ['Keratin\nprogramme'])
for t, f in zip(ax.get_xticklabels(), feats):
    if f != SCORE: t.set_fontstyle('italic')
ax.set_xlim(-0.6, xpos[SCORE] + 0.6)
ax.set_ylim(min(-0.2, meta.lo.min() - 0.1), max(ytop + [0.4]))
ax.set_ylabel("pooled Hedges' g (PTK2B high vs low)")
ax.set_title('Keratin programme with and without proliferation adjustment', loc='left')
hB = [Line2D([0], [0], marker='D', ls='', markerfacecolor=RED, markeredgecolor=RED, mew=1.6, markersize=8, label='Adjusted: stroma + immune'),
      Line2D([0], [0], marker='D', ls='', markerfacecolor='white', markeredgecolor=RED, mew=1.6, markersize=8, label='Adjusted: stroma + immune + proliferation')]
ax.legend(handles=hB, frameon=False, loc='upper center', bbox_to_anchor=(0.5, -0.14), ncol=2, fontsize=10)
fig.text(0.0, -0.05,
         'Five HGSOC cohorts; top vs bottom quartile of PTK2B; random-effects pooling; BH-FDR: * < 0.05, ** < 0.01, *** < 0.001. Proliferation: G1/S + G2/M genes (Tirosh et al. 2016).',
         ha='left', fontsize=9.5)
plt.tight_layout(); save_fig(fig, 'Fig3_PYK2_keratin_proliferation_adjusted')

CORDER = {v: i for i, v in enumerate(list(LAB.values()) + ['Pooled'])}
FORDER = {f: i for i, f in enumerate(KER + [SCORE])}

S1 = pd.DataFrame(cohort_rows)
tot = {'Cohort': 'Total', 'Reference': '', 'Accession': '', 'Platform': '', 'Data source': '', 'Median age, years': ''}
for c in ['Samples in source', 'HGSOC after inclusion criteria', 'Tumours analysed',
          'PTK2B-high (top quartile)', 'PTK2B-low (bottom quartile)']:
    tot[c] = int(pd.to_numeric(S1[c], errors='coerce').sum())
S1 = pd.concat([S1, pd.DataFrame([tot])], ignore_index=True)

S2 = pd.DataFrame(step_rows)

S3 = pd.DataFrame({
    'Gene': hit.index, 'Genome-wide rank': hit['rank'], 'Genes ranked': N,
    'Top % of genome': (100 * hit['rank'] / N).round(1),
    'Pooled z': hit.z, "Pooled Hedges' g": hit.pooled_g,
    '95% CI lower': hit.pooled_g - 1.96 * hit.se, '95% CI upper': hit.pooled_g + 1.96 * hit.se,
    'Cohorts (k)': hit.k, 'I² (%)': hit.I2, 'P value': hit.p, 'FDR (BH, genome-wide)': hit.FDR})

def src_table(model_list):
    out = []
    for model in model_list:
        p_, m_ = per[per.model == model], meta[meta.model == model]
        se_ = np.sqrt(p_.var_g)
        out.append(pd.DataFrame({
            'Feature': p_.feature, 'Adjustment': MN[model], 'Cohort': p_.cohort.map(LAB),
            'n PTK2B-high': p_.n_high, 'n PTK2B-low': p_.n_low,
            "Hedges' g": p_.g, '95% CI lower': p_.g - 1.96 * se_, '95% CI upper': p_.g + 1.96 * se_}))
        out.append(pd.DataFrame({
            'Feature': m_.feature, 'Adjustment': MN[model], 'Cohort': 'Pooled',
            'n PTK2B-high': m_.n_high, 'n PTK2B-low': m_.n_low,
            "Hedges' g": m_.pooled_g, '95% CI lower': m_.lo, '95% CI upper': m_.hi,
            'Cohorts (k)': m_.k, 'I² (%)': m_.I2, 'P value': m_.p, 'FDR (BH)': m_.FDR}))
    t = pd.concat(out, ignore_index=True)
    return t.assign(
        _f=t.Feature.map(FORDER),
        _m=t.Adjustment.map({v: i for i, v in enumerate(MN.values())}),
        _c=t.Cohort.map(CORDER)
    ).sort_values(['_f', '_m', '_c']).drop(columns=['_f', '_m', '_c'])

S4  = src_table(['adjusted']).drop(columns='Adjustment')
S5  = src_table(['adjusted', 'adjusted + proliferation'])
piv = meta.pivot(index='feature', columns='model', values='pooled_g')
sig = meta[meta.model == 'adjusted'].set_index('feature').FDR < 0.05
ret = (100 * piv['adjusted + proliferation'] / piv['adjusted']).round(0).where(sig)
S5['Effect retained after proliferation adjustment (%)'] = np.where(
    (S5.Cohort == 'Pooled') & (S5.Adjustment == MN['adjusted + proliferation']),
    S5.Feature.map(ret), np.nan)

defs = pd.DataFrame([
    {'Cohort': 'Definition', 'Gene set': 'Keratin programme (PYK2-dependent in HGS2 cells)',
     'Genes defined': 6, 'Genes passing expression filter': '', 'Genes used': ', '.join(KER)},
    {'Cohort': 'Definition', 'Gene set': 'Stromal score markers',
     'Genes defined': 10, 'Genes passing expression filter': '', 'Genes used': ', '.join(STROMAL)},
    {'Cohort': 'Definition', 'Gene set': 'Immune score markers',
     'Genes defined': 10, 'Genes passing expression filter': '', 'Genes used': ', '.join(IMMUNE)},
    {'Cohort': 'Definition', 'Gene set': 'Proliferation: G1/S phase genes (Tirosh et al. 2016)',
     'Genes defined': len(G1S), 'Genes passing expression filter': '', 'Genes used': ', '.join(G1S)},
    {'Cohort': 'Definition', 'Gene set': 'Proliferation: G2/M phase genes (Tirosh et al. 2016)',
     'Genes defined': len(G2M), 'Genes passing expression filter': '', 'Genes used': ', '.join(G2M)}])
S6 = pd.concat([defs, pd.DataFrame(used_rows)], ignore_index=True)

sheets = [
    ('S1_Cohorts', S1, 'Cohort characteristics: source, platform, tumours included and analysed, PTK2B groups.'),
    ('S2_Inclusion_steps', S2, 'Tumours remaining after each inclusion step, per cohort.'),
    ('S3_Keratin_genome_rank', S3, 'Figure 1: position of the six keratin-programme genes in the genome-wide ranking.'),
    ('S4_Keratin_per_dataset', S4, "Figure 2: keratin genes and keratin programme score, Hedges' g per cohort and pooled."),
    ('S5_Keratin_proliferation', S5, 'Figure 3: keratin genes and programme score with and without proliferation adjustment.'),
    ('S6_Gene_sets', S6, 'Gene-set definitions and genes passing the expression filter in each cohort.'),
]
if INCLUDE_GENOMEWIDE:
    sheets.append(('S7_Genomewide_ranking', pd.DataFrame({
        'Rank': de['rank'], 'Gene': de.index, "Pooled Hedges' g": de.pooled_g,
        'SE': de.se, 'Pooled z': de.z, 'Cohorts (k)': de.k, 'I² (%)': de.I2,
        'P value': de.p, 'FDR (BH)': de.FDR}),
        'All genes ranked by pooled z, PTK2B high vs low (source data for Figure 1).'))

readme = pd.DataFrame(
    [{'Sheet': n, 'Content': d} for n, _, d in sheets] +
    [{'Sheet': '', 'Content': ''},
     {'Sheet': 'Methods',
      'Content': 'Five independent HGSOC cohorts analysed separately and pooled by DerSimonian–Laird random-effects meta-analysis. PTK2B-high = top quartile, PTK2B-low = bottom quartile of PTK2B within each cohort. Within-cohort ranks residualised on stromal and immune marker scores (PTK2B residualised before grouping); in S5 the proliferation score was added as a further covariate. Effect size: Hedges\' g. Multiple testing: Benjamini–Hochberg FDR (genome-wide in S3; within the seven keratin tests per adjustment in S4–S5). Genes used only if expressed above the 25th percentile of mean expression in at least 10% of tumours; mitochondrial and ribosomal genes excluded. The keratin programme was defined in advance from PYK2-dependent genes in HGS2 cells.'},
     {'Sheet': 'Abbreviations',
      'Content': "g, Hedges' g (standardised mean difference); CI, confidence interval; FDR, false discovery rate; I², between-cohort heterogeneity; k, number of cohorts pooled."}])

with pd.ExcelWriter(XLSX, engine='openpyxl') as xw:
    readme.to_excel(xw, sheet_name='README', index=False)
    for n, df_, _ in sheets:
        df_.to_excel(xw, sheet_name=n, index=False)

wb   = load_workbook(XLSX)
thin = Side(style='thin', color='000000')
for ws in wb.worksheets:
    hdr = [c.value for c in ws[1]]
    for c in ws[1]:
        c.font      = Font(name='Arial', size=10, bold=True)
        c.fill      = PatternFill('solid', fgColor='E7E6E6')
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border    = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        bold = any(c.value in ('Pooled', 'Total') for c in row)
        for c in row:
            c.font = Font(name='Arial', size=10, bold=bold)
            h = str(hdr[c.column - 1] or '')
            if isinstance(c.value, float):
                c.number_format = (
                    '0.00E+00' if re.search(r'P value|FDR', h)
                    else '0'   if 'retained' in h
                    else '0.0' if re.search(r'I²|Top %|age', h, re.I)
                    else '0.000')
    for i, h in enumerate(hdr, 1):
        w = max([len(str(h or ''))] + [len(str(ws.cell(r, i).value or '')) for r in range(2, min(ws.max_row, 300) + 1)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 45)
    ws.freeze_panes      = 'A2'
    ws.row_dimensions[1].height = 42
ws = wb['README']
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 120
for r in range(2, ws.max_row + 1):
    ws.cell(r, 1).font      = Font(name='Arial', size=10, bold=True)
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
wb.save(XLSX)

try:
    from google.colab import files
    import shutil, tempfile
    tmp = os.path.join(tempfile.gettempdir(), 'Supplementary_Tables_PYK2_keratin_HGSOC.xlsx')
    shutil.copy(XLSX, tmp)
    files.download(tmp)
except Exception:
    pass

pd.set_option('display.width', 250); pd.set_option('display.max_columns', 20)
print('\n=== S1 cohorts ===')
print(S1.drop(columns=['Reference', 'Data source']).to_string(index=False))
print('\n=== Figure 1: keratin genes in the genome-wide ranking ===')
print(S3[['Gene', 'Genome-wide rank', 'Genes ranked', 'Top % of genome', 'Pooled z', 'FDR (BH, genome-wide)']].round(4).to_string(index=False))
print('\n=== Figures 2–3: pooled keratin results ===')
print(meta.pivot_table(index='feature', columns='model', values=['pooled_g', 'FDR']).round(3).to_string())
if SCORE in ret.index:
    print(f'\nKeratin programme: {ret[SCORE]:.0f}% of effect retained after proliferation adjustment')
print('\nAll figures and tables saved to:', OUT)
