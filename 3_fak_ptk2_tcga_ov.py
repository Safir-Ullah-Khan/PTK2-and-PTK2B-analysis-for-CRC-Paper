#!/usr/bin/env python3
import subprocess, sys
for pkg in ['openpyxl', 'statsmodels']:
    try:
        __import__(pkg)
    except ImportError:
        subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', pkg], check=True)

import os, re, gzip, urllib.request, warnings, textwrap
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from scipy import stats
from statsmodels.stats.multitest import multipletests
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

try:
    from google.colab import drive
    drive.mount('/content/drive')
except Exception:
    pass

BASE     = '/content/drive/MyDrive/FAK_PYK2_revision'
DATA_DIR = os.path.join(BASE, 'data')
os.makedirs(DATA_DIR, exist_ok=True)

HUB = 'https://gdc-hub.s3.us-east-1.amazonaws.com/download/'
SOURCES = {
    'counts':   ('TCGA-OV.star_counts.tsv.gz',
                 [HUB + 'TCGA-OV.star_counts.tsv.gz']),
    'probemap': ('gencode.v36.annotation.gtf.gene.probemap',
                 [HUB + 'gencode.v36.annotation.gtf.gene.probemap']),
}

SETS = {
    'Mesenchymal':        ['CDH2','VIM','FN1','SNAI1','SNAI2','ZEB1','ZEB2','TWIST1','TWIST2'],
    'Fibrillar collagen': ['COL1A1','COL1A2','COL3A1','COL5A1','COL5A2'],
}
SHORT = {'Mesenchymal': 'MES', 'Fibrillar collagen': 'COL'}

ROW = dict(
    label   = 'PTK2 (FAK)',
    pc      = 1,
    scores  = ['Mesenchymal', 'Fibrillar collagen'],
    scatter = 'Mesenchymal',
    markers = ['Mesenchymal', 'Fibrillar collagen'],
    tcol    = ['#FDDBC7', '#EF8A62', '#B2182B'],
    acc     = '#B2182B',
)
TERTS = ['Low', 'Mid', 'High']
BLACK = '#000000'; SEED = 1; N_BOOT = 5000

plt.rcParams.update({
    'font.family': 'sans-serif',
    'font.sans-serif': ['Arial', 'Liberation Sans', 'DejaVu Sans'],
    'font.size': 11, 'font.weight': 'bold',
    'axes.titlesize': 11, 'axes.titleweight': 'bold',
    'axes.labelsize': 11, 'axes.labelweight': 'bold',
    'xtick.labelsize': 10, 'ytick.labelsize': 10,
    'text.color': BLACK, 'axes.labelcolor': BLACK,
    'axes.edgecolor': BLACK, 'xtick.color': BLACK, 'ytick.color': BLACK,
    'axes.linewidth': 1.0,
    'axes.spines.top': False, 'axes.spines.right': False,
    'pdf.fonttype': 42, 'ps.fonttype': 42,
    'svg.fonttype': 'none', 'figure.dpi': 110,
})

def fetch(key):
    fname, urls = SOURCES[key]
    path = os.path.join(DATA_DIR, fname)
    if os.path.exists(path) and os.path.getsize(path) > 1000:
        print(f'{key}: using cached file'); return path
    print(f'{key}: downloading...')
    for u in urls:
        try:
            req = urllib.request.Request(u, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=900) as r, open(path, 'wb') as out:
                total = int(r.headers.get('Content-Length', 0))
                downloaded = 0
                while True:
                    chunk = r.read(1024*1024)
                    if not chunk:
                        break
                    out.write(chunk); downloaded += len(chunk)
                    if total:
                        print(f'  {downloaded/1e6:.0f} / {total/1e6:.0f} MB', end='\r')
            print(f'\n{key}: done'); return path
        except Exception as e:
            print(f'{key}: failed {u} ({type(e).__name__}: {e})')
    raise FileNotFoundError(
        f'Could not download {fname}.\n'
        f'Please manually place it in {DATA_DIR}\n'
        f'Direct URL: {urls[0]}'
    )

def read_tab(path):
    opener = gzip.open if path.endswith('.gz') else open
    with opener(path, 'rt') as fh:
        return pd.read_csv(fh, sep='\t', index_col=0, low_memory=False)

def fmt_p(p):
    if p < 1e-300: return '< 1e-300'
    if p < 0.001:
        m, e = f'{p:.1e}'.split('e')
        return f'{m} x 10^{int(e)}'
    return f'{p:.3f}' if p < 0.1 else f'{p:.2f}'

def cliffs(a, b):
    a, b = np.asarray(a), np.asarray(b)
    return ((a[:,None] > b[None,:]).sum() - (a[:,None] < b[None,:]).sum()) / (len(a)*len(b))

def pearson_ci(x, y):
    r, p = stats.pearsonr(x, y)
    z  = np.arctanh(r); se = 1/np.sqrt(len(x)-3)
    return r, np.tanh(z - 1.96*se), np.tanh(z + 1.96*se), p

pm  = pd.read_csv(fetch('probemap'), sep='\t', index_col=0)
cnt = read_tab(fetch('counts'))

if cnt.values.max() < 50:
    cnt = np.round(2**cnt - 1)

sym = pm['gene'].to_dict()
cnt.index = [sym.get(i, i) for i in cnt.index]
cnt = cnt.groupby(level=0).sum().clip(lower=0)
cnt = cnt.loc[[g for g in cnt.index
               if not re.match(r'^(MT-|RPL\d|RPS\d|MRPL\d|MRPS\d)', str(g))]]
cnt.columns = [str(c) for c in cnt.columns]
cnt = cnt[[c for c in cnt.columns if c[13:15] == '01']]
cnt = cnt.loc[:, ~pd.Index([c[:12] for c in cnt.columns]).duplicated()]
cnt.columns = [c[:12] for c in cnt.columns]
pts = list(cnt.columns); n = len(pts)

need    = ['PTK2'] + [g for v in SETS.values() for g in v]
missing = [g for g in need if g not in cnt.index]
if missing:
    raise SystemExit(f'Genes not found in count matrix: {missing}')
print(f'TCGA-OV primary tumours: {n}')

def log2v(genes, pc):
    x = cnt.loc[genes].astype(float)
    if pc == 0 and (x <= 0).any().any():
        print(f'  zeros present — using log2(value + 1)'); pc = 1
    return np.log2(x + pc)

def zscore(L):
    return L.sub(L.mean(1), axis=0).div(L.std(1, ddof=1), axis=0)

rng = np.random.default_rng(SEED)
R   = ROW
genes_all = [g for s in R['markers'] for g in SETS[s]]
L   = log2v(['PTK2'] + genes_all, R['pc'])
Z   = zscore(L.loc[genes_all])
k   = L.loc['PTK2'].values

c1, c2 = np.quantile(k, [1/3, 2/3])
tv = np.where(k <= c1, 'Low', np.where(k <= c2, 'Mid', 'High'))
SC = {s: Z.loc[SETS[s]].mean(0).values for s in R['scores']}

score_rows = []
for s, y in SC.items():
    g = [y[tv == t] for t in TERTS]
    boot = [rng.choice(g[2], len(g[2])).mean() - rng.choice(g[0], len(g[0])).mean()
            for _ in range(N_BOOT)]
    r, lo, hi, pr = pearson_ci(k, y)
    rho, prho = stats.spearmanr(k, y)
    ols = stats.linregress(k, y)
    mw_ml = stats.mannwhitneyu(g[1], g[0]).pvalue
    mw_hm = stats.mannwhitneyu(g[2], g[1]).pvalue
    score_rows.append({
        'Score': s,
        'Gene set (n genes)': f'{s} ({len(SETS[s])})',
        'n total': n,
        'Transform (PTK2)': f'log2(count + {R["pc"]})',
        'Tertile cutoff 1 (log2)': round(c1, 4),
        'Tertile cutoff 2 (log2)': round(c2, 4),
        'n Low': len(g[0]), 'n Mid': len(g[1]), 'n High': len(g[2]),
        'Mean Low': round(g[0].mean(), 4),
        'Mean Mid': round(g[1].mean(), 4),
        'Mean High': round(g[2].mean(), 4),
        'Median Low': round(float(np.median(g[0])), 4),
        'Median Mid': round(float(np.median(g[1])), 4),
        'Median High': round(float(np.median(g[2])), 4),
        'SD Low': round(g[0].std(ddof=1), 4),
        'SD Mid': round(g[1].std(ddof=1), 4),
        'SD High': round(g[2].std(ddof=1), 4),
        'IQR Low': round(float(np.percentile(g[0],75)-np.percentile(g[0],25)), 4),
        'IQR Mid': round(float(np.percentile(g[1],75)-np.percentile(g[1],25)), 4),
        'IQR High': round(float(np.percentile(g[2],75)-np.percentile(g[2],25)), 4),
        'Kruskal-Wallis H': round(stats.kruskal(*g).statistic, 4),
        'Kruskal-Wallis P': stats.kruskal(*g).pvalue,
        'MW Mid vs Low P': mw_ml,
        'MW High vs Mid P': mw_hm,
        'MW High vs Low P': stats.mannwhitneyu(g[2], g[0]).pvalue,
        "Cliff's delta High vs Low": round(cliffs(g[2], g[0]), 4),
        'High minus Low (mean diff)': round(g[2].mean() - g[0].mean(), 4),
        'Bootstrap mean diff (N=5000)': round(np.mean(boot), 4),
        'Bootstrap 95% CI lower': round(float(np.percentile(boot, 2.5)), 4),
        'Bootstrap 95% CI upper': round(float(np.percentile(boot, 97.5)), 4),
        'Pearson r': round(r, 4),
        'Pearson 95% CI lower': round(lo, 4),
        'Pearson 95% CI upper': round(hi, 4),
        'Pearson P': pr,
        'Spearman rho': round(rho, 4),
        'Spearman P': prho,
        'OLS slope': round(ols.slope, 4),
        'OLS intercept': round(ols.intercept, 4),
        'OLS R²': round(ols.rvalue**2, 4),
    })

tmp = []
for s in R['markers']:
    for gname in SETS[s]:
        r, lo, hi, pr = pearson_ci(k, L.loc[gname].values)
        rho, prho = stats.spearmanr(k, L.loc[gname].values)
        y_g = L.loc[gname].values
        grps = [y_g[tv == t] for t in TERTS]
        tmp.append({
            'Gene set': s, 'Gene': gname, 'n': n,
            'Mean Low': round(grps[0].mean(), 4),
            'Mean Mid': round(grps[1].mean(), 4),
            'Mean High': round(grps[2].mean(), 4),
            'MW High vs Low P': stats.mannwhitneyu(grps[2], grps[0]).pvalue,
            "Cliff's delta H vs L": round(cliffs(grps[2], grps[0]), 4),
            'Pearson r': round(r, 4),
            'Pearson CI lower': round(lo, 4),
            'Pearson CI upper': round(hi, 4),
            'Pearson P': pr,
            'Spearman rho': round(rho, 4),
            'Spearman P': prho,
        })
tmp = pd.DataFrame(tmp)
tmp['Pearson BH q']  = multipletests(tmp['Pearson P'],  method='fdr_bh')[1]
tmp['Spearman BH q'] = multipletests(tmp['Spearman P'], method='fdr_bh')[1]
tmp['MW BH q']       = multipletests(tmp['MW High vs Low P'], method='fdr_bh')[1]
M_df = tmp.copy()
Mk   = tmp.set_index('Gene')
S_df = pd.DataFrame(score_rows)

def srow(s):
    return S_df[S_df.Score == s].iloc[0]

fig = plt.figure(figsize=(16.0, 5.5))
gsp = fig.add_gridspec(1, 4, width_ratios=[0.9, 1.25, 1.05, 0.75], wspace=0.48)
tcol, acc = R['tcol'], R['acc']

ax = fig.add_subplot(gsp[0, 0])
x  = k; y = SC[R['scatter']]
for t, c in zip(TERTS, tcol):
    m = tv == t
    ax.scatter(x[m], y[m], s=10, c=c, edgecolors=BLACK,
               linewidths=0.25, alpha=0.9, zorder=2)
X    = np.column_stack([np.ones_like(x), x])
beta = np.linalg.lstsq(X, y, rcond=None)[0]
xs   = np.linspace(x.min(), x.max(), 200)
Xs   = np.column_stack([np.ones_like(xs), xs])
res  = y - X @ beta; s2 = res @ res / (len(y)-2)
se   = np.sqrt(np.einsum('ij,jk,ik->i', Xs, s2*np.linalg.inv(X.T@X), Xs))
tc   = stats.t.ppf(0.975, len(y)-2)
ax.fill_between(xs, Xs@beta - tc*se, Xs@beta + tc*se, color=acc, alpha=0.18, lw=0, zorder=1)
ax.plot(xs, Xs@beta, color=acc, lw=2.2, zorder=3)
for v in (c1, c2):
    ax.axvline(v, color='#404040', lw=1.3, ls='--', zorder=0)
rr = srow(R['scatter'])
ax.text(0.03, 0.97,
        f'r = {rr["Pearson r"]:.2f}, P = {fmt_p(rr["Pearson P"])}\n'
        f'rho = {rr["Spearman rho"]:.2f}, P = {fmt_p(rr["Spearman P"])}',
        transform=ax.transAxes, va='top', fontsize=9.5, fontweight='bold', zorder=5,
        bbox=dict(facecolor='white', edgecolor='none', alpha=0.85, pad=1.5))
ylo, yhi = ax.get_ylim()
ax.set_ylim(ylo, yhi + 0.22*(yhi-ylo))
ax.set_xlabel('log2(PTK2 + 1)')
ax.set_ylabel(f'{R["scatter"]} score (mean z)')
ax.set_title('a  PTK2 vs mesenchymal score', loc='left')

sub   = gsp[0, 1].subgridspec(2, 2, height_ratios=[0.07, 1],
                                width_ratios=[1, 0.035], hspace=0.04, wspace=0.04)
order = np.argsort(x)
Zm    = Z.loc[genes_all].values[:, order]
axk   = fig.add_subplot(sub[0, 0])
axh   = fig.add_subplot(sub[1, 0])
axc   = fig.add_subplot(sub[1, 1])

kz = (x - x.mean()) / x.std(ddof=1)
axk.imshow(kz[order][None,:], aspect='auto', cmap='Reds', vmin=-2, vmax=2, interpolation='nearest')
axk.set_xticks([]); axk.set_yticks([0]); axk.set_yticklabels(['PTK2'], fontsize=10)
for sp in axk.spines.values(): sp.set_visible(False)

im   = axh.imshow(Zm, aspect='auto', cmap='RdBu_r', vmin=-3, vmax=3, interpolation='nearest')
cuts = np.searchsorted(np.sort(x), (c1, c2), side='right')
for c_ in cuts:
    axh.axvline(c_-0.5, color=BLACK, lw=1.2)
    axk.axvline(c_-0.5, color=BLACK, lw=1.2)
y0 = -0.5
for i, s in enumerate(R['markers']):
    if i: axh.axhline(y0, color=BLACK, lw=1.0)
    y0 += len(SETS[s])
axh.set_yticks(range(len(genes_all)))
axh.set_yticklabels(genes_all, style='italic', fontsize=9.5)
axh.set_xticks([cuts[0]/2, (cuts[0]+cuts[1])/2, (cuts[1]+len(order))/2])
axh.set_xticklabels(TERTS)
axh.tick_params(axis='x', length=0)
axh.set_xlabel('Tumours ordered by PTK2')
cb = fig.colorbar(im, cax=axc); cb.set_label('z-score', fontsize=10); cb.ax.tick_params(labelsize=9.5)
axk.set_title('b  Marker expression by PTK2', loc='left')

ax      = fig.add_subplot(gsp[0, 2])
centres = []
for j, s in enumerate(R['scores']):
    yv     = SC[s]
    groups = [yv[tv == t] for t in TERTS]
    pos    = j*4.8 + np.array([1, 2, 3])
    centres.append(pos[1])
    bp = ax.boxplot(groups, positions=pos, widths=0.62, patch_artist=True,
                    showfliers=False,
                    medianprops=dict(color=BLACK, lw=2.2),
                    whiskerprops=dict(lw=1.1), capprops=dict(lw=1.1),
                    boxprops=dict(lw=1.1), zorder=2)
    for pch, c in zip(bp['boxes'], tcol):
        pch.set_facecolor(c); pch.set_edgecolor(BLACK)
    for p_, gv in zip(pos, groups):
        ax.scatter(p_ + rng.uniform(-0.2, 0.2, len(gv)), gv,
                   s=4, color=BLACK, alpha=0.35, lw=0, zorder=3)
        ax.scatter(p_, gv.mean(), marker='D', s=26,
                   color='white', edgecolors=BLACK, linewidths=1.1, zorder=4)
    rr  = srow(s); top = max(np.max(g) for g in groups)
    ax.plot([pos[0], pos[0], pos[2], pos[2]],
            [top+0.15, top+0.3, top+0.3, top+0.15], color=BLACK, lw=1)
    dl = rr["Cliff's delta High vs Low"]
    ax.text(pos[1], top+0.38,
            f'KW P = {fmt_p(rr["Kruskal-Wallis P"])}\n'
            f'H vs L P = {fmt_p(rr["MW High vs Low P"])}\ndelta = {dl:.2f}',
            ha='center', va='bottom', fontsize=8.5, fontweight='bold', linespacing=1.2)

ax.set_xticks(centres)
ax.set_xticklabels([s.replace('Fibrillar collagen', 'Fibrillar\ncollagen') for s in R['scores']])
ax.tick_params(axis='x', length=0)
ax.axhline(0, color='#525252', lw=1.0, ls='--', zorder=0)
ax.set_xlim(-0.9, 4.8*(len(R['scores'])-1)+4.6)
ylo, yhi = ax.get_ylim(); ax.set_ylim(ylo, yhi+1.5)
ax.set_ylabel('Score (mean z)')
nt = [int((tv == t).sum()) for t in TERTS]
ax.legend(
    handles=[Patch(facecolor=c, edgecolor=BLACK, label=f'{t} ({k_})')
             for t, c, k_ in zip(TERTS, tcol, nt)] +
            [Line2D([], [], marker='D', ls='', mfc='white', mec=BLACK, ms=6, label='Mean')],
    fontsize=8.5, frameon=False, ncol=4, loc='upper center',
    bbox_to_anchor=(0.5, -0.15), handletextpad=0.3, columnspacing=0.6, handlelength=1.2)
ax.set_title('c  Scores by PTK2 tertile', loc='left')

ax = fig.add_subplot(gsp[0, 3])
for yi, gname in enumerate(genes_all):
    rv, lo, hi, q = Mk.loc[gname, ['Pearson r','Pearson CI lower','Pearson CI upper','Pearson BH q']]
    ax.barh(yi, rv, height=0.7, color=acc if q < 0.05 else '#D9D9D9',
            edgecolor=BLACK, lw=0.6, zorder=2)
    ax.plot([lo, hi], [yi, yi], color=BLACK, lw=1.2, zorder=3)
ax.axvline(0, color=BLACK, lw=0.8)
ax.set_yticks(range(len(genes_all)))
ax.set_yticklabels(genes_all, style='italic', fontsize=9.5)
ax.set_ylim(len(genes_all)-0.4, -0.6)
xr = max(0.6, np.ceil((np.nanmax(np.abs(Mk[['Pearson CI lower','Pearson CI upper']].values))+0.05)*10)/10)
ax.set_xlim(-0.2 if Mk['Pearson CI lower'].min() > -0.15 else -xr, xr)
y0 = -0.5
for i, s in enumerate(R['markers']):
    y1 = y0 + len(SETS[s])
    if i: ax.axhline(y0, color='#404040', lw=1.3, ls='--')
    ax.text(xr*1.04, (y0+y1)/2, SHORT[s], rotation=270, va='center', fontsize=10.5)
    y0 = y1
ax.set_xlabel('Pearson r (+/- 95% CI)')
ax.legend(handles=[Patch(facecolor=acc, edgecolor=BLACK, label='BH q < 0.05'),
                   Patch(facecolor='#D9D9D9', edgecolor=BLACK, label='BH q >= 0.05')],
          fontsize=8.5, frameon=False, ncol=2, loc='upper center',
          handlelength=1.2, bbox_to_anchor=(0.45, -0.18))
ax.set_title('d  Marker-level correlation', loc='left')

NOTE = (
    f'TCGA-OV primary tumours (GDC STAR counts), n = {n}. '
    'PTK2 (FAK): log2(value + 1). Scores: mean gene-wise z-scores (sample SD). '
    'Tertiles of PTK2 expression. Boxes: median/IQR, whiskers 1.5xIQR, diamond = mean. '
    "KW: Kruskal-Wallis; H vs L: Mann-Whitney high vs low; delta: Cliff's delta; "
    'q: Benjamini-Hochberg. Mesenchymal score: CDH2, VIM, FN1, SNAI1, SNAI2, ZEB1, ZEB2, TWIST1, TWIST2. '
    'Fibrillar collagen score: COL1A1, COL1A2, COL3A1, COL5A1, COL5A2 (Ojalill et al., Matrix Biol 2026).'
)
fig.text(0.01, 0.0, textwrap.fill(NOTE, 220), fontsize=8, ha='left', va='top')
plt.suptitle('PTK2 (FAK) — Mesenchymal & Fibrillar Collagen Programmes in TCGA-OV',
             fontsize=12, fontweight='bold', y=1.01)
plt.tight_layout(pad=1.5)
plt.show()

xlsx = '/tmp/FAK_PTK2_source_data.xlsx'

S1 = pd.DataFrame({
    'Patient':                  pts,
    'PTK2_raw_count':           cnt.loc['PTK2'].values.astype(int),
    'PTK2_log2_count_plus1':    k.round(6),
    'PTK2_tertile':             tv,
})
for s in R['scores']:
    S1[f'{s}_score_mean_z'] = SC[s].round(6)
for gname in genes_all:
    S1[f'{gname}_log2'] = L.loc[gname].values.round(6)
    S1[f'{gname}_z']    = Z.loc[gname].values.round(6)
S1 = S1.sort_values('PTK2_log2_count_plus1').reset_index(drop=True)

S2 = S_df.copy()

S3 = M_df[['Gene set','Gene','n',
            'Mean Low','Mean Mid','Mean High',
            'MW High vs Low P','MW BH q',
            "Cliff's delta H vs L",
            'Pearson r','Pearson CI lower','Pearson CI upper','Pearson P','Pearson BH q',
            'Spearman rho','Spearman P','Spearman BH q']].copy()

tert_rows = []
for s in R['scores']:
    y = SC[s]
    for t in TERTS:
        g = y[tv == t]
        tert_rows.append({
            'Score': s, 'Tertile': t, 'n': len(g),
            'Mean': round(g.mean(), 4), 'Median': round(float(np.median(g)), 4),
            'SD': round(g.std(ddof=1), 4),
            'IQR': round(float(np.percentile(g,75)-np.percentile(g,25)), 4),
            'Min': round(float(g.min()), 4), 'Max': round(float(g.max()), 4),
            'Q25': round(float(np.percentile(g, 25)), 4),
            'Q75': round(float(np.percentile(g, 75)), 4),
        })
for gname in genes_all:
    y = L.loc[gname].values
    for t in TERTS:
        g = y[tv == t]
        tert_rows.append({
            'Score': f'{gname} (log2)', 'Tertile': t, 'n': len(g),
            'Mean': round(g.mean(), 4), 'Median': round(float(np.median(g)), 4),
            'SD': round(g.std(ddof=1), 4),
            'IQR': round(float(np.percentile(g,75)-np.percentile(g,25)), 4),
            'Min': round(float(g.min()), 4), 'Max': round(float(g.max()), 4),
            'Q25': round(float(np.percentile(g, 25)), 4),
            'Q75': round(float(np.percentile(g, 75)), 4),
        })
S4 = pd.DataFrame(tert_rows)

S5 = pd.DataFrame([
    {'Gene set': 'Mesenchymal', 'n genes': len(SETS['Mesenchymal']),
     'Genes': ', '.join(SETS['Mesenchymal']),
     'Reference': 'Thiery et al., Cell 2009; standard EMT markers'},
    {'Gene set': 'Fibrillar collagen', 'n genes': len(SETS['Fibrillar collagen']),
     'Genes': ', '.join(SETS['Fibrillar collagen']),
     'Reference': 'Ojalill et al., Matrix Biol 2026'},
    {'Gene set': 'Excluded (noise)', 'n genes': '',
     'Genes': 'MT-* (mitochondrial), RPL*/RPS* (ribosomal), MRPL*/MRPS* (mito-ribosomal)',
     'Reference': 'Standard quality filter'},
])

S6 = cnt.loc[['PTK2'] + genes_all].T.reset_index().rename(columns={'index': 'Patient'})
S6.insert(1, 'PTK2_tertile', pd.Series(tv, index=range(len(tv))))
tert_map = dict(zip(pts, tv))
S6['PTK2_tertile'] = S6['Patient'].map(tert_map)

sheets = [
    ('README', pd.DataFrame({
        'Sheet': ['S1_Per_patient','S2_Score_statistics','S3_Marker_correlations',
                  'S4_Tertile_descriptives','S5_Gene_sets','S6_Raw_counts'],
        'Content': [
            f'All {n} TCGA-OV primary tumours: PTK2 raw count, log2, tertile assignment, both score values, and individual marker gene log2 and z-scores. Sorted by PTK2 log2.',
            'Score-level statistics for Mesenchymal and Fibrillar collagen scores: tertile cutoffs, group means/medians/SDs/IQRs, Kruskal-Wallis (with H statistic), Mann-Whitney (all three pairwise), Cliff\'s delta, bootstrap mean difference 95% CI (N=5,000), Pearson r with Fisher 95% CI, Spearman rho, OLS slope/intercept/R².',
            'Per-gene Pearson r (Fisher 95% CI), Spearman rho, Mann-Whitney High vs Low, and Cliff\'s delta for all 14 marker genes, with Benjamini-Hochberg FDR across all genes. Also includes mean log2 expression per tertile.',
            'Descriptive statistics (n, mean, median, SD, IQR, Q25, Q75, min, max) per tertile for each score and each individual marker gene (log2).',
            'Gene set definitions, member genes, and references for Mesenchymal and Fibrillar collagen signatures; exclusion criteria for noise genes.',
            'Raw GDC STAR read counts for PTK2 and all 14 marker genes across all primary tumours, with tertile annotation.',
        ],
    }), 'col'),
    ('S1_Per_patient',          S1,  'num'),
    ('S2_Score_statistics',     S2,  'num'),
    ('S3_Marker_correlations',  S3,  'num'),
    ('S4_Tertile_descriptives', S4,  'num'),
    ('S5_Gene_sets',            S5,  'str'),
    ('S6_Raw_counts',           S6,  'num'),
]

with pd.ExcelWriter(xlsx, engine='openpyxl') as w:
    for name, df_, _ in sheets:
        df_.to_excel(w, sheet_name=name, index=False)

wb  = load_workbook(xlsx)
thin = Side(style='thin', color='000000')
HDR  = PatternFill('solid', fgColor='E7E6E6')

for ws in wb.worksheets:
    hdr_vals = [c.value for c in ws[1]]
    for c in ws[1]:
        c.font      = Font(name='Arial', size=10, bold=True)
        c.fill      = HDR
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border    = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            h = str(hdr_vals[c.column - 1] or '')
            c.font = Font(name='Arial', size=10)
            if isinstance(c.value, float):
                if re.search(r'\bP\b|BH q', h):
                    c.number_format = '0.00E+00'
                elif re.search(r'R²|r =|rho|delta|slope|intercept|CI|mean|median|SD|IQR|score|log2|Min|Max|Q\d', h, re.I):
                    c.number_format = '0.0000'
                else:
                    c.number_format = '0.000'
    for i, h in enumerate(hdr_vals, 1):
        w_col = max(
            [len(str(h or ''))] +
            [len(str(ws.cell(r, i).value or '')) for r in range(2, min(ws.max_row, 200) + 1)]
        )
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w_col + 2), 48)
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 36

ws_rm = wb['README']
ws_rm.column_dimensions['A'].width = 26
ws_rm.column_dimensions['B'].width = 110
for r in range(2, ws_rm.max_row + 1):
    ws_rm.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
    ws_rm.row_dimensions[r].height = 52

wb.save(xlsx)

try:
    from google.colab import files
    files.download(xlsx)
    print('Excel download started — check your browser downloads.')
except Exception:
    print(f'Not in Colab — file written to: {xlsx}')

print(f'Excel:  FAK_PTK2_source_data.xlsx  (7 sheets: README + S1-S6)')
print('\n=== Score summary ===')
print(S_df[['Score','n total','n Low','n Mid','n High',
            'Mean Low','Mean High','Kruskal-Wallis P',
            'MW High vs Low P',"Cliff's delta High vs Low",
            'Pearson r','Spearman rho']].to_string(
    index=False, float_format=lambda v: f'{v:.3g}'))
